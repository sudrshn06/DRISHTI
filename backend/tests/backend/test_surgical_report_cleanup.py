import io
from pathlib import Path

import pypdf
import pytest
from docx import Document

from app.schemas.report import (
    CaptureSummary,
    CaptureViewSummary,
    DeclarationFindingItem,
    ExtractedEvidenceItem,
    InspectionReportSnapshot,
    InspectorContextSnapshot,
    OverallDisposition,
    ReportEvidenceAsset,
    ReportMetadata,
    ReportSummaryCounts,
    VisualComplianceFindingItem,
)
from app.services.docx_report_service import generate_docx_report
from app.services.pdf_report_service import generate_pdf_report
from app.services.report_service import (
    FSSAI_2026_REGULATORY_NOTE,
    compute_overall_disposition,
    present_finding_reason,
    resolve_evidence_asset,
    resolve_report_disposition,
)


PROJECT_ROOT = Path(__file__).resolve().parents[3]


def _source(relative_path: str) -> str:
    path = PROJECT_ROOT / relative_path
    if not path.exists():
        pytest.skip("frontend source is not mounted in the backend container")
    return path.read_text(encoding="utf-8")


def _review_finding() -> DeclarationFindingItem:
    return DeclarationFindingItem(
        rule_id="UNIT_SALE_PRICE_DECLARATION_PRESENCE",
        field="UNIT_SALE_PRICE",
        status="REVIEW_REQUIRED",
        reason="Evidence was not sufficient to confirm the declaration.",
        legal_reference="LMPC Rules 2011, Rule 6(11)",
        applicability_status="APPLICABLE",
    )


def test_complete_required_capture_is_not_presented_as_incomplete_for_absence_limit() -> None:
    disposition, reason = compute_overall_disposition(
        statutory_findings=[_review_finding()],
        visual_findings=[],
        capture_status="COMPLETE_EVIDENCE_CAPTURE",
        evidence_sufficiency="INSUFFICIENT_FOR_ABSENCE_EVALUATION",
    )

    assert disposition == OverallDisposition.REVIEW_REQUIRED
    assert "capture is incomplete" not in reason.lower()


def test_missing_required_capture_remains_incomplete() -> None:
    disposition, _ = compute_overall_disposition(
        statutory_findings=[_review_finding()],
        visual_findings=[],
        capture_status="INCOMPLETE_INSPECTION",
        evidence_sufficiency="INSUFFICIENT_FOR_ABSENCE_EVALUATION",
    )

    assert disposition == OverallDisposition.INCOMPLETE_INSPECTION


def _snapshot(*, capture_complete: bool = True) -> InspectionReportSnapshot:
    review = _review_finding()
    return InspectionReportSnapshot(
        metadata=ReportMetadata(
            report_id="report-surgical",
            inspection_id="inspection-surgical",
            generated_at="2026-08-29T12:00:00+00:00",
            reference_date="2026-08-29",
            product_category="GENERIC_RETAIL_PACKAGE",
            capture_plan_id="plan-surgical",
        ),
        # Deliberately preserve the legacy stored value. Renderers must project
        # it without mutating the immutable snapshot.
        overall_disposition=OverallDisposition.INCOMPLETE_INSPECTION,
        disposition_reason="Package surface capture is incomplete.",
        summary_counts=ReportSummaryCounts(
            total_statutory_checks=4,
            statutory_pass_count=2,
            statutory_fail_count=0,
            statutory_review_required_count=1,
            statutory_not_applicable_count=1,
            total_visual_checks=1,
            visual_review_required_count=0,
            visual_not_evaluable_count=1,
            visual_observation_clear_count=0,
        ),
        inspector_context=InspectorContextSnapshot(
            product_origin="DOMESTIC",
            regulatory_product_class="FOOD",
            date_regulatory_regime="FOOD",
            date_package_exemption="NONE",
            is_electronic="NON_ELECTRONIC",
            package_structure="SINGLE",
            alcohol_context="NON_ALCOHOLIC",
        ),
        capture_summary=CaptureSummary(
            views=[
                CaptureViewSummary(
                    view_id="FRONT",
                    display_name="Front",
                    required=True,
                    captured=True,
                ),
                CaptureViewSummary(
                    view_id="BACK",
                    display_name="Back",
                    required=True,
                    captured=capture_complete,
                ),
                CaptureViewSummary(
                    view_id="SIDE_LEFT",
                    display_name="Left side",
                    required=False,
                    captured=False,
                ),
            ],
            total_views_required=2,
            captured_required_count=2 if capture_complete else 1,
            missing_required_views=[] if capture_complete else ["BACK"],
            capture_status="COMPLETE_EVIDENCE_CAPTURE" if capture_complete else "INCOMPLETE_INSPECTION",
            evidence_sufficiency="INSUFFICIENT_FOR_ABSENCE_EVALUATION",
            overall_quality_status="ACCEPTABLE",
        ),
        declaration_findings=[review],
        visual_compliance_findings=[
            VisualComplianceFindingItem(
                rule_id="RULE_7_MINIMUM_NUMERAL_HEIGHT",
                field="NET_QUANTITY",
                capability="NOT_EVALUABLE_FROM_CURRENT_CAPTURE",
                status="NOT_EVALUABLE",
                reason="Physical scale is unavailable.",
                legal_reference="LMPC Rule 7",
                limitations="Calibrated measurement is required.",
            )
        ],
        extracted_evidence=[
            ExtractedEvidenceItem(
                field="MRP",
                status="DETECTED",
                raw_value="MRP ₹: 110.00 (incl. of all taxes)",
                normalized_value={"currency": "₹", "amount": 110.0},
                confidence=1.0,
                view_ids=["BACK"],
            )
        ],
    )


def _pdf_text(data: bytes) -> str:
    return "\n".join(page.extract_text() or "" for page in pypdf.PdfReader(io.BytesIO(data)).pages)


def _docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    values = [p.text for p in doc.paragraphs]
    values.extend(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    return "\n".join(values)


def test_legacy_snapshot_is_projected_consistently_without_mutation() -> None:
    snapshot = _snapshot()

    disposition, _ = resolve_report_disposition(snapshot)
    pdf_text = _pdf_text(generate_pdf_report(snapshot))
    docx_text = _docx_text(generate_docx_report(snapshot))

    assert disposition == OverallDisposition.REVIEW_REQUIRED
    assert snapshot.overall_disposition == OverallDisposition.INCOMPLETE_INSPECTION
    assert "OVERALL DISPOSITION: REVIEW REQUIRED" in pdf_text
    assert "OVERALL DISPOSITION: REVIEW REQUIRED" in docx_text
    assert "OVERALL DISPOSITION: INCOMPLETE INSPECTION" not in pdf_text
    assert "OVERALL DISPOSITION: INCOMPLETE INSPECTION" not in docx_text


def test_legacy_snapshot_with_missing_required_view_still_renders_incomplete() -> None:
    snapshot = _snapshot(capture_complete=False)

    assert resolve_report_disposition(snapshot)[0] == OverallDisposition.INCOMPLETE_INSPECTION
    assert "OVERALL DISPOSITION: INCOMPLETE INSPECTION" in _pdf_text(generate_pdf_report(snapshot))
    assert "OVERALL DISPOSITION: INCOMPLETE INSPECTION" in _docx_text(generate_docx_report(snapshot))


def test_statutory_counts_reconcile() -> None:
    counts = _snapshot().summary_counts
    displayed_categories = (
        counts.statutory_pass_count
        + counts.statutory_fail_count
        + counts.statutory_review_required_count
        + counts.statutory_not_applicable_count
    )
    assert displayed_categories == counts.total_statutory_checks


def test_country_origin_reason_separates_package_evidence_from_officer_context() -> None:
    finding = DeclarationFindingItem(
        rule_id="COUNTRY_OF_ORIGIN_DECLARATION_PRESENCE",
        field="COUNTRY_OF_ORIGIN",
        status="NOT_APPLICABLE",
        reason="Explicit origin is DOMESTIC.",
        legal_reference="LMPC Rule 6(1)(aa)",
        applicability_status="NOT_APPLICABLE",
    )
    reason = present_finding_reason(finding, _snapshot().inspector_context)

    assert "No explicit country-of-origin declaration was detected in the package evidence" in reason
    assert "officer-provided product-origin classification: Domestic" in reason
    assert "Explicit origin is DOMESTIC" not in reason


def test_fssai_wording_is_precise_and_shared_note_is_not_repeated_per_row() -> None:
    veg = DeclarationFindingItem(
        rule_id="FSSAI_VEG_NONVEG_SYMBOL",
        field="FSSAI_VEG_NONVEG",
        status="REVIEW_REQUIRED",
        reason=f"Supporting text 'None' was detected. ({FSSAI_2026_REGULATORY_NOTE})",
        legal_reference="FSSAI Regulation 5(4)",
        applicability_status="APPLICABLE",
    )
    result = present_finding_reason(veg)

    assert "Supporting text 'None'" not in result
    assert "No reliable visual veg/non-veg symbol evidence" in result
    assert FSSAI_2026_REGULATORY_NOTE not in result


def test_evidence_assets_map_to_final_field_evidence_ids() -> None:
    report = _snapshot()
    report.evidence_assets = [
        ReportEvidenceAsset(capture_id="front", view_id="FRONT", image_sha256="a" * 64, evidence_ids=["net"]),
        ReportEvidenceAsset(capture_id="back", view_id="BACK", image_sha256="b" * 64, evidence_ids=["mrp", "nutrition"]),
    ]
    fields = {
        "MRP": "mrp",
        "NET_QUANTITY": "net",
        "FSSAI_NUTRITION": "nutrition",
    }
    expected_capture = {"MRP": "back", "NET_QUANTITY": "front", "FSSAI_NUTRITION": "back"}

    for field, evidence_id in fields.items():
        item = ExtractedEvidenceItem(
            field=field,
            status="DETECTED",
            evidence_ids=[evidence_id],
            capture_ids=["front", "back"],
            has_geometry=True,
            pixel_box={"x_min": 1, "y_min": 1, "x_max": 10, "y_max": 10},
        )
        asset, exact_region, _ = resolve_evidence_asset(report, item)
        assert asset.capture_id == expected_capture[field]
        assert exact_region is True


def test_gemini_only_context_never_receives_a_fake_exact_region() -> None:
    report = _snapshot()
    report.evidence_assets = [
        ReportEvidenceAsset(capture_id="back", view_id="BACK", image_sha256="b" * 64),
    ]
    item = ExtractedEvidenceItem(
        field="MONTH_YEAR",
        status="DETECTED",
        raw_value="Use By: 14/05/2025",
        evidence_ids=[],
        capture_ids=["back"],
        has_geometry=True,
        pixel_box={"x_min": 1, "y_min": 1, "x_max": 10, "y_max": 10},
    )

    asset, exact_region, use_polygons = resolve_evidence_asset(report, item)
    assert asset.capture_id == "back"
    assert exact_region is False
    assert use_polygons is False


def test_rupee_and_officer_readable_report_labels_are_preserved() -> None:
    from app.services.docx_report_service import _format_normalized_value as docx_value
    from app.services.pdf_report_service import _format_normalized_value as pdf_value

    value = {"currency": "₹", "amount": 110.0}
    assert pdf_value(value) == "₹110.00"
    assert docx_value(value) == "₹110.00"
    assert "Final Detected Value" in _pdf_text(generate_pdf_report(_snapshot()))
    assert "Final Detected Value" in _docx_text(generate_docx_report(_snapshot()))
    assert "Origin View" not in _pdf_text(generate_pdf_report(_snapshot()))
    assert "Origin View" not in _docx_text(generate_docx_report(_snapshot()))


def test_finalized_dashboard_has_no_active_correction_or_incomplete_wording() -> None:
    summary = _source("frontend/src/components/inspection/OfficerResultSummary.jsx")
    package_card = _source("frontend/src/components/inspection/PackageInformationCard.jsx")
    inspection = _source("frontend/src/pages/MultiViewInspection.jsx")
    history = _source("frontend/src/pages/InspectionHistory.jsx")
    dashboard = _source("frontend/src/pages/Dashboard.jsx")

    assert "Finalized with documented review items" in summary
    assert "They are not confirmed non-compliance findings" in summary
    assert "!isFinalized && (" in package_card
    assert "Package information confirmed" in package_card
    assert "guidance && !isFinalized" in inspection
    for overview in (history, dashboard):
        assert "item.lifecycle_status === 'FINALIZED' && item.overall_disposition === 'INCOMPLETE_INSPECTION'" in overview
        assert "Documented Review Items" in overview


def test_dashboard_counts_roles_context_and_complaint_wording_are_explicit() -> None:
    inspection = _source("frontend/src/pages/MultiViewInspection.jsx")
    package_card = _source("frontend/src/components/inspection/PackageInformationCard.jsx")
    evidence_panel = _source("frontend/src/components/ocr/EvidenceInspectorPanel.jsx")

    for category in ("Compliant", "Non-compliant", "Review", "Not applicable"):
        assert category in inspection
    assert "statutoryFindings.length" in inspection
    assert "regulatoryEscalationAvailable &&" in inspection
    assert "Nothing is submitted automatically." in inspection
    assert "Officer-provided Product Origin Classification" in inspection
    assert "Business declaration" in package_card
    assert "Business Declaration" in evidence_panel
    assert "Product Origin Classification" in evidence_panel


def test_visual_actions_and_confidence_labels_are_not_misleading() -> None:
    evidence_panel = _source("frontend/src/components/ocr/EvidenceInspectorPanel.jsx")
    ocr_view = _source("frontend/src/components/ocr/OCRResultView.jsx")

    assert "Calibrated Measurement Required" in evidence_panel
    assert "Physical Verification Required" in evidence_panel
    assert "Controlled-Lighting Verification Required" in evidence_panel
    assert "if (visualRule.status === 'NEEDS_RECAPTURE') return 'Take Another Photograph'" in evidence_panel
    assert "Machine Confidence" in ocr_view
