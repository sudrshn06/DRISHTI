import io
import json
import hashlib
import numpy as np
import cv2
import pytest
from datetime import date
from fastapi.testclient import TestClient
from docx import Document
import pypdf

from app.main import app
from app.schemas.ocr import FieldCandidate
from app.schemas.inspection import (
    InspectionSession,
    CaptureRecord,
    CapturePlan,
    CaptureViewRequirement
)
from app.schemas.visual_assessment import (
    CaptureVisualAssessmentSummary,
    DeclarationVisualAssessment,
    RegionGeometry,
    PixelBoundingBox,
    NormalizedBoundingBox,
    VisualReadabilitySignal,
    VisualCheckCapability,
    VisualObservationStatus,
    VisualRuleEvaluationResult
)
from app.schemas.compliance import LegalStatus, RuleEvaluationResult
from app.schemas.applicability import ApplicabilityDecision, ApplicabilityStatus
from app.schemas.report import OverallDisposition, InspectionReportSnapshot
from app.services.report_service import generate_inspection_report
from app.services.image_store import store_capture_image, clear_image_store
from app.services.pdf_report_service import generate_pdf_report
from app.services.docx_report_service import generate_docx_report

from app.api.deps import get_current_user, get_current_user_optional
from app.db.session import SessionLocal
from app.models.user import UserModel

client = TestClient(app)


@pytest.fixture(scope="module", autouse=True)
def persist_mock_user():
    """The authenticated principal must exist when PostgreSQL enforces ownership FKs."""
    with SessionLocal() as db:
        if db.get(UserModel, "test-inspector-uuid-001") is None:
            db.add(UserModel(
                user_id="test-inspector-uuid-001",
                username="test_inspector",
                email="inspector@drishti.local",
                password_hash="hash",
                full_name="Test Inspector",
                role="INSPECTOR",
                is_active=True,
            ))
            db.commit()
    yield

@pytest.fixture(autouse=True)
def override_auth():
    mock_user = UserModel(
        user_id="test-inspector-uuid-001",
        username="test_inspector",
        email="inspector@drishti.local",
        password_hash="hash",
        full_name="Test Inspector",
        role="INSPECTOR",
        is_active=True
    )
    app.dependency_overrides[get_current_user] = lambda: mock_user
    app.dependency_overrides[get_current_user_optional] = lambda: mock_user
    yield
    app.dependency_overrides.pop(get_current_user, None)
    app.dependency_overrides.pop(get_current_user_optional, None)

TEST_PLAN = CapturePlan(
    capture_plan_id="plan_software_1",
    name="Standard Software Inspection",
    views=[
        CaptureViewRequirement(view_id="FRONT", display_name="Front View", required=True),
        CaptureViewRequirement(view_id="BACK", display_name="Back View", required=True)
    ],
    absence_evaluation_eligible=True
)

def create_synthetic_image(width=600, height=400) -> bytes:
    img = np.full((height, width, 3), (240, 240, 240), dtype=np.uint8)
    cv2.putText(img, "TEST COMMODITY", (50, 150), cv2.FONT_HERSHEY_SIMPLEX, 1.0, (20, 20, 20), 2)
    success, encoded = cv2.imencode(".jpg", img)
    return encoded.tobytes()

def extract_pdf_text(pdf_bytes: bytes) -> str:
    reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)

def extract_docx_text(docx_bytes: bytes) -> str:
    doc = Document(io.BytesIO(docx_bytes))
    full_text = []
    for p in doc.paragraphs:
        if p.text:
            full_text.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                full_text.append(" | ".join(row_text))
    return "\n".join(full_text)

def test_1_cross_format_consistency_json_pdf_docx():
    # 1. Verify JSON snapshot, PDF, and DOCX produce matching substantive results
    img_bytes = create_synthetic_image()
    sha = hashlib.sha256(img_bytes).hexdigest()
    store_capture_image("cap_6e_01", sha, img_bytes, "image/jpeg", 600, 400)
    
    cand_mrp = FieldCandidate(
        field="MRP",
        status="DETECTED",
        raw_value="MRP Rs. 350.00",
        normalized_value={"currency": "INR", "amount": 350.0},
        confidence=0.97,
        evidence_ids=["ev_6e_mrp"],
        capture_ids=["cap_6e_01"]
    )
    
    geom = RegionGeometry(
        pixel_box=PixelBoundingBox(x_min=50.0, y_min=100.0, x_max=250.0, y_max=150.0, width_px=200.0, height_px=50.0, area_px=10000.0),
        normalized_box=NormalizedBoundingBox(x_min=0.083, y_min=0.25, x_max=0.417, y_max=0.375),
        image_width=600,
        image_height=400,
        width_ratio=0.33,
        height_ratio=0.125,
        area_ratio=0.042,
        distance_to_edge_px={"top": 100.0, "bottom": 250.0, "left": 50.0, "right": 350.0},
        touches_edge=False
    )
    
    decl_vis = DeclarationVisualAssessment(
        field="MRP",
        evidence_ids=["ev_6e_mrp"],
        capture_id="cap_6e_01",
        view_id="FRONT",
        raw_text="MRP Rs. 350.00",
        polygons=[[[50.0, 100.0], [250.0, 100.0], [250.0, 150.0], [50.0, 150.0]]],
        readability=VisualReadabilitySignal(
            ocr_confidence=0.97,
            is_clipped=False,
            capture_quality_status="ACCEPTABLE",
            technical_readability="EVALUABLE"
        ),
        technical_status="EVALUABLE",
        geometry=geom
    )
    
    rule_mrp = RuleEvaluationResult(
        rule_id="MRP_DECLARATION_PRESENCE",
        field="MRP",
        rule_version="1.0",
        reference_date=date(2025, 1, 1),
        status=LegalStatus.PASS,
        reason="Maximum Retail Price declared with currency symbol and amount.",
        source_reference="Legal Metrology Rules, 2011 Rule 6(1)(e)",
        evidence_ids=["ev_6e_mrp"],
        applicability=ApplicabilityDecision(rule_id="MRP_DECLARATION_PRESENCE", status=ApplicabilityStatus.APPLICABLE, is_applicable=True, reason="Standard commodity")
    )
    
    session = InspectionSession(
        inspection_id="insp_6e_consistency",
        reference_date="2025-01-01",
        product_category="RETAIL COMMODITY",
        product_origin="DOMESTIC",
        regulatory_product_class="NON_FOOD",
        date_regulatory_regime="GENERAL",
        date_package_exemption="NONE",
        is_electronic="NON_ELECTRONIC",
        package_structure="SINGLE",
        alcohol_context="NON_ALCOHOLIC",
        capture_plan_id="plan_software_1",
        capture_status="COMPLETE_EVIDENCE_CAPTURE",
        evidence_sufficiency="SUFFICIENT_FOR_ABSENCE_EVALUATION",
        aggregated_candidates=[cand_mrp],
        rule_evaluations=[rule_mrp],
        visual_rule_evaluations=[],
        captures=[
            CaptureRecord(
                capture_id="cap_6e_01",
                view_id="FRONT",
                evidence_id="ev_6e_mrp",
                image_sha256=sha,
                visual_assessment=CaptureVisualAssessmentSummary(
                    capture_id="cap_6e_01",
                    view_id="FRONT",
                    image_width=600,
                    image_height=400,
                    assessments=[decl_vis]
                ),
                status="ACCEPTED",
                pipeline_status="COMPLETED"
            )
        ]
    )
    
    # Generate Snapshot
    snapshot = generate_inspection_report(session, TEST_PLAN)
    
    # 1. JSON Snapshot assertions
    assert snapshot.metadata.inspection_id == "insp_6e_consistency"
    assert snapshot.overall_disposition == OverallDisposition.NO_VIOLATIONS_DETECTED_IN_EVALUATED_SCOPE
    assert len(snapshot.declaration_findings) == 1
    assert snapshot.declaration_findings[0].status == "PASS"
    assert snapshot.declaration_findings[0].legal_reference == "Legal Metrology Rules, 2011 Rule 6(1)(e)"
    assert snapshot.evidence_assets[0].image_sha256 == sha
    
    # 2. PDF assertions
    pdf_bytes = generate_pdf_report(snapshot)
    pdf_text = extract_pdf_text(pdf_bytes)
    assert snapshot.metadata.report_id in pdf_text
    assert "insp_6e_consistency" in pdf_text
    assert "NO VIOLATIONS DETECTED" in pdf_text
    assert "Rule 6(1)(e)" in pdf_text
    assert "ev_6e_mrp" in pdf_text
    assert sha[:16] in pdf_text
    assert "Compliance Score" not in pdf_text
    
    # 3. DOCX assertions
    docx_bytes = generate_docx_report(snapshot)
    docx_text = extract_docx_text(docx_bytes)
    assert snapshot.metadata.report_id in docx_text
    assert "insp_6e_consistency" in docx_text
    assert "NO VIOLATIONS DETECTED" in docx_text
    assert "Rule 6(1)(e)" in docx_text
    assert "ev_6e_mrp" in docx_text
    assert sha[:16] in docx_text
    assert "Compliance Score" not in docx_text

def test_2_snapshot_immutability_against_session_mutation():
    # 2. Mutating active session context does not mutate existing snapshot
    session = InspectionSession(
        inspection_id="insp_6e_immutability",
        reference_date="2025-01-01",
        product_category="RETAIL COMMODITY",
        product_origin="DOMESTIC",
        regulatory_product_class="NON_FOOD",
        capture_plan_id="plan_software_1",
        capture_status="COMPLETE_EVIDENCE_CAPTURE",
        evidence_sufficiency="SUFFICIENT_FOR_ABSENCE_EVALUATION"
    )
    
    snapshot = generate_inspection_report(session, TEST_PLAN)
    initial_origin = snapshot.inspector_context.product_origin
    assert initial_origin == "DOMESTIC"
    
    # Mutate session
    session.product_origin = "IMPORTED"
    session.regulatory_product_class = "FOOD"
    
    # Snapshot remains unchanged
    assert snapshot.inspector_context.product_origin == "DOMESTIC"
    assert snapshot.inspector_context.regulatory_product_class == "NON_FOOD"

def test_3_api_report_endpoints_and_error_handling():
    # 3. Test POST/GET JSON, PDF, and DOCX endpoints via FastAPI TestClient
    init_res = client.post("/api/inspections", data={
        "reference_date": "2025-01-01",
        "product_category": "ELECTRONICS",
        "capture_plan_id": "plan_software_1"
    })
    assert init_res.status_code == 200
    inspection_id = init_res.json()["inspection_id"]
    
    # POST /api/inspections/{id}/report
    post_rep = client.post(f"/api/inspections/{inspection_id}/report")
    assert post_rep.status_code == 200
    rep_data = post_rep.json()
    assert "metadata" in rep_data
    assert "overall_disposition" in rep_data
    report_id = rep_data["metadata"]["report_id"]
    
    # GET /api/inspections/{id}/report
    get_rep = client.get(f"/api/inspections/{inspection_id}/report")
    assert get_rep.status_code == 200
    assert get_rep.json()["metadata"]["report_id"] == report_id
    
    # GET /api/inspections/{id}/report.pdf
    pdf_res = client.get(f"/api/inspections/{inspection_id}/report.pdf")
    assert pdf_res.status_code == 200
    assert pdf_res.headers["content-type"] == "application/pdf"
    assert f'filename="DRISHTI_Inspection_Report_{report_id}.pdf"' in pdf_res.headers["content-disposition"]
    assert len(pdf_res.content) > 1000
    
    # GET /api/inspections/{id}/report.docx
    docx_res = client.get(f"/api/inspections/{inspection_id}/report.docx")
    assert docx_res.status_code == 200
    assert docx_res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert f'filename="DRISHTI_Inspection_Report_{report_id}.docx"' in docx_res.headers["content-disposition"]
    assert len(docx_res.content) > 1000
    
    # 404 on non-existent inspection
    assert client.get("/api/inspections/non_existent_id/report").status_code == 404
    assert client.get("/api/inspections/non_existent_id/report.pdf").status_code == 404
    assert client.get("/api/inspections/non_existent_id/report.docx").status_code == 404

def test_4_all_disposition_invariants_and_no_scores():
    # 4. Assert strict enum dispositions and total absence of compliance scores
    for disp in [
        OverallDisposition.VIOLATIONS_FOUND,
        OverallDisposition.INCOMPLETE_INSPECTION,
        OverallDisposition.REVIEW_REQUIRED,
        OverallDisposition.NO_VIOLATIONS_DETECTED_IN_EVALUATED_SCOPE
    ]:
        assert disp.value in ("VIOLATIONS_FOUND", "INCOMPLETE_INSPECTION", "REVIEW_REQUIRED", "NO_VIOLATIONS_DETECTED_IN_EVALUATED_SCOPE")
        assert disp.value not in ("COMPLIANT", "FULLY_COMPLIANT", "APPROVED", "FAIL")
