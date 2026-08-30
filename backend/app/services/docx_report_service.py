import io
from typing import Optional, Dict, Any, List
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from app.schemas.report import (
    InspectionReportSnapshot,
    OverallDisposition,
    DeclarationFindingItem,
    VisualComplianceFindingItem,
    ExtractedEvidenceItem
)
from app.services.report_service import (
    FSSAI_2026_REGULATORY_NOTE,
    present_finding_reason,
    resolve_evidence_asset,
    resolve_report_disposition,
)

def set_cell_background(cell, fill_hex: str):
    """Sets the background fill color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal cell padding in dxa (1 pt = 20 dxa)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """Sets uniform subtle borders on a table."""
    tbl_pr = table._tbl.tblPr
    tbl_borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tbl_pr.append(tbl_borders)


def set_repeat_table_header(row):
    """Repeat the first row when a Word table continues on another page."""
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

def _format_context_value(val: Optional[str]) -> str:
    if not val or val == "UNKNOWN":
        return "Unknown / Not established"
    return val.replace("_", " ")

def _format_normalized_value(val: Optional[Dict[str, Any]]) -> str:
    if not val:
        return "—"
    if "currency" in val and "amount" in val:
        currency = "₹" if val.get("currency") in ("INR", "₹") else str(val.get("currency") or "")
        amount = val.get("amount")
        amount_text = f"{amount:.2f}" if isinstance(amount, (int, float)) else str(amount)
        if "per_unit" in val:
            per_q = val.get("per_quantity", 1.0)
            q_str = f"{per_q} " if per_q != 1.0 else ""
            return f"{currency}{amount_text} / {q_str}{val['per_unit']}"
        return f"{currency}{amount_text}"
    if "value" in val and "unit" in val:
        amount = val["value"]
        amount_text = f"{amount:g}" if isinstance(amount, (int, float)) else str(amount)
        return f"{amount_text} {val['unit']}"
    if "role" in val or "name" in val:
        parts = []
        if val.get("role"):
            parts.append(f"Role: {str(val['role']).replace('_', ' ').title()}")
        if val.get("name"):
            parts.append(f"Entity: {val['name']}")
        if val.get("address"):
            parts.append(f"Address: {val['address']}")
        if val.get("pin_code"):
            parts.append(f"PIN: {val['pin_code']}")
        return " - ".join(parts)
    if "country_text" in val:
        return f"{val.get('declaration_type', 'ORIGIN')}: {val['country_text']}"
    if "name_text" in val:
        return val["name_text"]
    if "month" in val and "year" in val:
        d_type = val.get("type", "DATE")
        day_str = f"{val['day']}/" if val.get("day") else ""
        return f"{d_type} {day_str}{val['month']:02d}/{val['year']}"
    if "email" in val or "phone" in val:
        parts = []
        if val.get("email"):
            parts.append(f"Email: {val['email']}")
        if val.get("phone"):
            parts.append(f"Phone: {val['phone']}")
        return ", ".join(parts)
    return ", ".join(f"{k}: {v}" for k, v in val.items() if v is not None)


def _declaration_label(ev: ExtractedEvidenceItem) -> str:
    if ev.field == "MANUFACTURER_PACKER_IMPORTER":
        return "Business Declaration"
    labels = {
        "BRAND_NAME": "Brand",
        "COMMON_GENERIC_NAME": "Product",
        "MRP": "MRP",
        "NET_QUANTITY": "Net Quantity",
        "MONTH_YEAR": "Date Declaration",
        "COUNTRY_OF_ORIGIN": "Country of Origin",
        "CONSUMER_CARE": "Consumer Care",
        "FSSAI_LICENCE": "FSSAI Licence",
        "FSSAI_INGREDIENTS": "Ingredients",
        "FSSAI_ALLERGENS": "Allergens",
        "FSSAI_NUTRITION": "Nutrition",
        "FSSAI_VEG_NONVEG": "Veg / Non-Veg Mark",
    }
    return labels.get(ev.field, ev.field.replace("_", " ").title())


def _final_detected_value(ev: ExtractedEvidenceItem) -> str:
    normalized = _format_normalized_value(ev.normalized_value)
    return normalized if normalized != "—" else (ev.raw_value or "—")

def generate_docx_report(report: InspectionReportSnapshot) -> bytes:
    """
    Renders an immutable InspectionReportSnapshot into a clean, professional,
    fully-editable Word (.docx) inspection report using python-docx.
    """
    doc = Document()

    # Configure 0.5-inch margins
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.text = f"DRISHTI — Packaged Commodity Legal Metrology Inspection Report | Report ID: {report.metadata.report_id}"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.runs[0].font.name = "Calibri"
        hp.runs[0].font.size = Pt(8)
        hp.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "Generated by DRISHTI — decision-support inspection software. Not a statutory judicial order."
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fp.runs[0].font.name = "Calibri"
        fp.runs[0].font.size = Pt(8)
        fp.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Set normal style font
    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Calibri'
    font_normal.size = Pt(9.5)
    font_normal.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    # =========================================================================
    # 1. HEADER & COVER SECTION
    # =========================================================================
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("DRISHTI: PACKAGED COMMODITY INSPECTION REPORT")
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(16)
    title_run.bold = True
    title_run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    title_p.paragraph_format.space_after = Pt(2)

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run("Legal Metrology (Packaged Commodities) Rules, 2011 — Decision Support & Evidence Record")
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(9.5)
    sub_run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
    sub_p.paragraph_format.space_after = Pt(8)

    # Metadata Grid Table
    meta = report.metadata
    meta_table = doc.add_table(rows=3, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(meta_table)

    col_widths = [Inches(1.3), Inches(2.45), Inches(1.3), Inches(2.45)]
    meta_rows = [
        [("Report ID:", True), (meta.report_id, False), ("Reference Date:", True), (meta.reference_date, False)],
        [("Inspection ID:", True), (meta.inspection_id, False), ("Generated (UTC):", True), (meta.generated_at[:19].replace("T", " "), False)],
        [("Product Category:", True), (meta.product_category.replace("_", " "), False), ("Capture Plan:", True), (meta.capture_plan_name or meta.capture_plan_id, False)]
    ]

    for r_idx, row_data in enumerate(meta_rows):
        row = meta_table.rows[r_idx]
        for c_idx, (text_val, is_bold) in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            run = p.add_run(text_val)
            run.font.name = "Calibri"
            run.font.size = Pt(8.5)
            run.bold = is_bold
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A) if is_bold else RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Overall Disposition Banner Table
    disp, presented_reason = resolve_report_disposition(report)
    disp_text = str(disp.value if hasattr(disp, "value") else disp).replace("_", " ")

    if disp == OverallDisposition.VIOLATIONS_FOUND:
        disp_bg = "FEF2F2"
        border_col = "DC2626"
        disp_r, disp_g, disp_b = 0x99, 0x1B, 0x1B
    elif disp == OverallDisposition.INCOMPLETE_INSPECTION:
        disp_bg = "FFFBEB"
        border_col = "D97706"
        disp_r, disp_g, disp_b = 0x92, 0x40, 0x0E
    elif disp == OverallDisposition.REVIEW_REQUIRED:
        disp_bg = "FEF3C7"
        border_col = "F59E0B"
        disp_r, disp_g, disp_b = 0xB4, 0x53, 0x09
    else: # NO_VIOLATIONS_DETECTED_IN_EVALUATED_SCOPE
        disp_bg = "ECFDF5"
        border_col = "059669"
        disp_r, disp_g, disp_b = 0x06, 0x5F, 0x46

    disp_table = doc.add_table(rows=1, cols=1)
    disp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(disp_table, color=border_col, sz="12")
    disp_cell = disp_table.rows[0].cells[0]
    disp_cell.width = Inches(7.5)
    set_cell_background(disp_cell, disp_bg)
    set_cell_margins(disp_cell, top=140, bottom=140, left=180, right=180)

    disp_p = disp_cell.paragraphs[0]
    disp_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    disp_p.paragraph_format.space_after = Pt(2)
    disp_run = disp_p.add_run(f"OVERALL DISPOSITION: {disp_text}")
    disp_run.font.name = "Calibri"
    disp_run.font.size = Pt(12)
    disp_run.bold = True
    disp_run.font.color.rgb = RGBColor(disp_r, disp_g, disp_b)

    reason_p = disp_cell.add_paragraph()
    reason_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    reason_p.paragraph_format.space_after = Pt(0)
    reason_lead = reason_p.add_run("Assessment Summary: ")
    reason_lead.bold = True
    reason_lead.font.size = Pt(9)
    reason_body = reason_p.add_run(presented_reason)
    reason_body.font.size = Pt(9)
    reason_body.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # =========================================================================
    # 2. SUMMARY COUNTS (DETERMINISTIC CATEGORICAL METRICS ONLY)
    # =========================================================================
    h1 = doc.add_heading("1. Compliance & Inspection Metrics", level=2)
    h1.paragraph_format.space_before = Pt(8)
    h1.paragraph_format.space_after = Pt(4)
    h1.paragraph_format.keep_with_next = True
    h1.runs[0].font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    counts = report.summary_counts

    # Statutory Table
    stat_table = doc.add_table(rows=2, cols=5)
    stat_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(stat_table)
    set_repeat_table_header(stat_table.rows[0])

    stat_col_widths = [Inches(2.5), Inches(1.25), Inches(1.25), Inches(1.25), Inches(1.25)]
    stat_headers = ["Statutory Checks Evaluated", "PASS", "FAIL", "REVIEW REQUIRED", "NOT APPLICABLE"]
    stat_values = [
        str(counts.total_statutory_checks),
        str(counts.statutory_pass_count),
        str(counts.statutory_fail_count),
        str(counts.statutory_review_required_count),
        str(counts.statutory_not_applicable_count)
    ]
    stat_colors = [
        RGBColor(0x0F, 0x17, 0x2A),
        RGBColor(0x05, 0x96, 0x69),
        RGBColor(0xDC, 0x26, 0x26),
        RGBColor(0xD9, 0x77, 0x06),
        RGBColor(0x64, 0x74, 0x8B)
    ]

    for c_idx, h_text in enumerate(stat_headers):
        cell = stat_table.rows[0].cells[c_idx]
        cell.width = stat_col_widths[c_idx]
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h_text)
        run.font.size = Pt(8.5)
        run.bold = True
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    for c_idx, val_text in enumerate(stat_values):
        cell = stat_table.rows[1].cells[c_idx]
        cell.width = stat_col_widths[c_idx]
        set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(val_text)
        run.font.size = Pt(8.5)
        run.bold = (c_idx in [1, 2, 3])
        run.font.color.rgb = stat_colors[c_idx]

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # Visual Table
    vis_table = doc.add_table(rows=2, cols=4)
    vis_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(vis_table)
    set_repeat_table_header(vis_table.rows[0])

    vis_col_widths = [Inches(2.5), Inches(1.66), Inches(1.66), Inches(1.68)]
    vis_headers = ["Visual Checks (Rule 7, 8, 9)", "Clear Observation", "Inspector Review Req.", "Not Evaluable from 2D Photo"]
    vis_values = [
        str(counts.total_visual_checks),
        str(counts.visual_observation_clear_count),
        str(counts.visual_review_required_count),
        str(counts.visual_not_evaluable_count)
    ]
    vis_colors = [
        RGBColor(0x0F, 0x17, 0x2A),
        RGBColor(0x05, 0x96, 0x69),
        RGBColor(0xD9, 0x77, 0x06),
        RGBColor(0x64, 0x74, 0x8B)
    ]

    for c_idx, h_text in enumerate(vis_headers):
        cell = vis_table.rows[0].cells[c_idx]
        cell.width = vis_col_widths[c_idx]
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h_text)
        run.font.size = Pt(8.5)
        run.bold = True
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    for c_idx, val_text in enumerate(vis_values):
        cell = vis_table.rows[1].cells[c_idx]
        cell.width = vis_col_widths[c_idx]
        set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(val_text)
        run.font.size = Pt(8.5)
        run.font.color.rgb = vis_colors[c_idx]

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # =========================================================================
    # 3. SURFACE CAPTURE & EVIDENCE SUFFICIENCY
    # =========================================================================
    h2 = doc.add_heading("2. Surface Capture & Evidence Sufficiency", level=2)
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.keep_with_next = True
    h2.runs[0].font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    c_sum = report.capture_summary
    c_p = doc.add_paragraph()
    c_p.paragraph_format.space_after = Pt(4)
    c_p.add_run(f"Capture Status: ").bold = True
    c_p.add_run(f"{c_sum.capture_status.replace('_', ' ')}  |  ")
    c_p.add_run(f"Evidence Sufficiency: ").bold = True
    c_p.add_run(f"{c_sum.evidence_sufficiency.replace('_', ' ')}  |  ")
    c_p.add_run(f"Overall Quality: ").bold = True
    c_p.add_run(f"{c_sum.overall_quality_status}  |  ")
    c_p.add_run(f"Required Views Captured: ").bold = True
    c_p.add_run(f"{c_sum.captured_required_count} of {c_sum.total_views_required}")

    view_table = doc.add_table(rows=len(c_sum.views) + 1, cols=5)
    view_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(view_table)
    set_repeat_table_header(view_table.rows[0])

    v_widths = [Inches(1.5), Inches(0.9), Inches(0.9), Inches(1.2), Inches(3.0)]
    v_headers = ["View Role", "Requirement", "Status", "Capture ID", "Image Quality / SHA-256"]

    for c_idx, h_text in enumerate(v_headers):
        cell = view_table.rows[0].cells[c_idx]
        cell.width = v_widths[c_idx]
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h_text)
        run.font.size = Pt(8.5)
        run.bold = True
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    for r_idx, v in enumerate(c_sum.views):
        row = view_table.rows[r_idx + 1]
        req_str = "Required" if v.required else "Optional"
        stat_str = "Captured" if v.captured else "Missing"
        cap_id_str = f"{v.capture_id[:12]}..." if v.capture_id else "—"

        q_notes = []
        if v.quality_status:
            q_notes.append(f"Quality: {v.quality_status}")
        if v.reasons:
            q_notes.extend(v.reasons)
        if v.image_sha256:
            q_notes.append(f"SHA: {v.image_sha256[:16]}...")
        q_str = "; ".join(q_notes) if q_notes else "Acceptable"

        row_vals = [f"{v.display_name} ({v.view_id})", req_str, stat_str, cap_id_str, q_str]
        for c_idx, text_val in enumerate(row_vals):
            cell = row.cells[c_idx]
            cell.width = v_widths[c_idx]
            set_cell_margins(cell, top=50, bottom=50, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text_val)
            run.font.size = Pt(8.5)
            if c_idx == 2:
                run.bold = True
                run.font.color.rgb = RGBColor(0x05, 0x96, 0x69) if v.captured else RGBColor(0xDC, 0x26, 0x26)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # =========================================================================
    # 4. STATUTORY DECLARATION FINDINGS
    # =========================================================================
    h3 = doc.add_heading("3A. Legal Metrology Findings", level=2)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True
    h3.runs[0].font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    decl_table = doc.add_table(rows=len(report.declaration_findings) + 1, cols=5)
    decl_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(decl_table)
    set_repeat_table_header(decl_table.rows[0])

    d_widths = [Inches(1.6), Inches(1.2), Inches(1.0), Inches(2.7), Inches(1.0)]
    d_headers = ["Rule & Declaration", "Legal Ref", "Status", "Statutory Justification / Findings", "Evidence"]

    for c_idx, h_text in enumerate(d_headers):
        cell = decl_table.rows[0].cells[c_idx]
        cell.width = d_widths[c_idx]
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h_text)
        run.font.size = Pt(8.5)
        run.bold = True
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    for r_idx, f in enumerate(report.declaration_findings):
        row = decl_table.rows[r_idx + 1]
        ev_str = ", ".join(eid[:8] for eid in f.evidence_ids) if f.evidence_ids else "—"

        # Col 0: Field & Rule ID
        p0 = row.cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0_lead = p0.add_run(f"{f.field.replace('_', ' ')}\n")
        r0_lead.bold = True
        r0_lead.font.size = Pt(8.5)
        r0_sub = p0.add_run(f.rule_id)
        r0_sub.font.size = Pt(7.5)
        r0_sub.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

        # Col 1: Legal Ref
        p1 = row.cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(f.legal_reference)
        r1.font.size = Pt(8.5)

        # Col 2: Status
        p2 = row.cells[2].paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(f.status)
        r2.bold = True
        r2.font.size = Pt(8.5)
        if f.status == "PASS":
            r2.font.color.rgb = RGBColor(0x05, 0x96, 0x69)
        elif f.status == "FAIL":
            r2.font.color.rgb = RGBColor(0xDC, 0x26, 0x26)
        elif f.status == "REVIEW_REQUIRED":
            r2.font.color.rgb = RGBColor(0xD9, 0x77, 0x06)
        else:
            r2.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

        # Col 3: Reason
        p3 = row.cells[3].paragraphs[0]
        p3.paragraph_format.space_after = Pt(0)
        r3 = p3.add_run(present_finding_reason(f, report.inspector_context))
        r3.font.size = Pt(8.5)

        # Col 4: Evidence IDs
        p4 = row.cells[4].paragraphs[0]
        p4.paragraph_format.space_after = Pt(0)
        r4 = p4.add_run(ev_str)
        r4.font.name = "Courier New"
        r4.font.size = Pt(7.5)
        r4.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

        for c_idx in range(5):
            row.cells[c_idx].width = d_widths[c_idx]
            set_cell_margins(row.cells[c_idx], top=50, bottom=50, left=100, right=100)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    if getattr(report, "food_label_findings", None):
        h3b = doc.add_heading("3B. Food Labelling / FSSAI Findings", level=2)
        h3b.paragraph_format.space_before = Pt(8)
        h3b.paragraph_format.space_after = Pt(4)
        h3b.paragraph_format.keep_with_next = True
        h3b.runs[0].font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

        food_table = doc.add_table(rows=len(report.food_label_findings) + 1, cols=5)
        food_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(food_table)
        set_repeat_table_header(food_table.rows[0])

        for c_idx, h_text in enumerate(d_headers):
            cell = food_table.rows[0].cells[c_idx]
            cell.width = d_widths[c_idx]
            set_cell_background(cell, "F1F5F9")
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(h_text)
            run.font.size = Pt(8.5)
            run.bold = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

        for r_idx, f in enumerate(report.food_label_findings):
            row = food_table.rows[r_idx + 1]
            ev_str = ", ".join(eid[:8] for eid in f.evidence_ids) if f.evidence_ids else "—"

            # Col 0: Field & Rule ID
            p0 = row.cells[0].paragraphs[0]
            p0.paragraph_format.space_after = Pt(0)
            r0_lead = p0.add_run(f"{f.field.replace('_', ' ')}\n")
            r0_lead.bold = True
            r0_lead.font.size = Pt(8.5)
            r0_sub = p0.add_run(f.rule_id)
            r0_sub.font.size = Pt(7.5)
            r0_sub.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

            # Col 1: Legal Ref
            p1 = row.cells[1].paragraphs[0]
            p1.paragraph_format.space_after = Pt(0)
            r1 = p1.add_run(f.legal_reference)
            r1.font.size = Pt(8.5)

            # Col 2: Status
            p2 = row.cells[2].paragraphs[0]
            p2.paragraph_format.space_after = Pt(0)
            r2 = p2.add_run(f.status)
            r2.bold = True
            r2.font.size = Pt(8.5)
            if f.status == "PASS":
                r2.font.color.rgb = RGBColor(0x05, 0x96, 0x69)
            elif f.status == "FAIL":
                r2.font.color.rgb = RGBColor(0xDC, 0x26, 0x26)
            elif f.status == "REVIEW_REQUIRED":
                r2.font.color.rgb = RGBColor(0xD9, 0x77, 0x06)
            else:
                r2.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

            # Col 3: Reason
            p3 = row.cells[3].paragraphs[0]
            p3.paragraph_format.space_after = Pt(0)
            r3 = p3.add_run(present_finding_reason(f, report.inspector_context))
            r3.font.size = Pt(8.5)

            # Col 4: Evidence IDs
            p4 = row.cells[4].paragraphs[0]
            p4.paragraph_format.space_after = Pt(0)
            r4 = p4.add_run(ev_str)
            r4.font.name = "Courier New"
            r4.font.size = Pt(7.5)
            r4.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

            for c_idx in range(5):
                row.cells[c_idx].width = d_widths[c_idx]
                set_cell_margins(row.cells[c_idx], top=50, bottom=50, left=100, right=100)

        regulatory_note = doc.add_paragraph()
        regulatory_note.paragraph_format.space_after = Pt(4)
        regulatory_note.add_run("Regulatory Notes: ").bold = True
        regulatory_note.add_run(f"{FSSAI_2026_REGULATORY_NOTE}.")
        for run in regulatory_note.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    # =========================================================================
    # 5. VISUAL COMPLIANCE FINDINGS (RULE 7, 8, 9)
    # =========================================================================
    h4 = doc.add_heading("4. Visual Compliance Observations (Rule 7, 8, 9)", level=2)
    h4.paragraph_format.space_before = Pt(8)
    h4.paragraph_format.space_after = Pt(2)
    h4.paragraph_format.keep_with_next = True
    h4.runs[0].font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    v_note = doc.add_paragraph()
    v_note.paragraph_format.space_after = Pt(4)
    v_note_run = v_note.add_run(
        "Note: Visual observations reflect 2D image analysis. Technical limitations (e.g. lack of physical scale for Rule 7) "
        "are explicitly marked NOT EVALUABLE and are never converted into false statutory non-compliance."
    )
    v_note_run.italic = True
    v_note_run.font.size = Pt(8.5)
    v_note_run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    vis_obs_table = doc.add_table(rows=len(report.visual_compliance_findings) + 1, cols=4)
    vis_obs_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(vis_obs_table)
    set_repeat_table_header(vis_obs_table.rows[0])

    vo_widths = [Inches(1.8), Inches(1.4), Inches(1.3), Inches(3.0)]
    vo_headers = ["Visual Check", "Capability", "Observation Status", "Observation & Statutory Limitations"]

    for c_idx, h_text in enumerate(vo_headers):
        cell = vis_obs_table.rows[0].cells[c_idx]
        cell.width = vo_widths[c_idx]
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h_text)
        run.font.size = Pt(8.5)
        run.bold = True
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    for r_idx, vr in enumerate(report.visual_compliance_findings):
        row = vis_obs_table.rows[r_idx + 1]

        # Col 0: Check Name & Ref
        p0 = row.cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0_lead = p0.add_run(f"{vr.rule_id.replace('_', ' ')}\n")
        r0_lead.bold = True
        r0_lead.font.size = Pt(8.5)
        r0_sub = p0.add_run(vr.legal_reference)
        r0_sub.font.size = Pt(7.5)
        r0_sub.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

        # Col 1: Capability
        p1 = row.cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(vr.capability.replace("_", " "))
        r1.font.size = Pt(8.5)

        # Col 2: Status
        p2 = row.cells[2].paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(vr.status.replace("_", " "))
        r2.bold = True
        r2.font.size = Pt(8.5)
        if vr.status == "OBSERVATION_CLEAR":
            r2.font.color.rgb = RGBColor(0x05, 0x96, 0x69)
        elif vr.status == "REVIEW_REQUIRED":
            r2.font.color.rgb = RGBColor(0xD9, 0x77, 0x06)
        elif vr.status == "NEEDS_RECAPTURE":
            r2.font.color.rgb = RGBColor(0xEA, 0x58, 0x0C)
        else:
            r2.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

        # Col 3: Observation & Limitations
        p3 = row.cells[3].paragraphs[0]
        p3.paragraph_format.space_after = Pt(0)
        r3_lead = p3.add_run("Observation: ")
        r3_lead.bold = True
        r3_lead.font.size = Pt(8.5)
        r3_body = p3.add_run(f"{vr.reason}\n")
        r3_body.font.size = Pt(8.5)

        if vr.limitations:
            r3_lim_lead = p3.add_run("Statutory Limitations: ")
            r3_lim_lead.bold = True
            r3_lim_lead.font.size = Pt(7.5)
            r3_lim_lead.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
            r3_lim_body = p3.add_run(vr.limitations)
            r3_lim_body.font.size = Pt(7.5)
            r3_lim_body.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

        for c_idx in range(4):
            row.cells[c_idx].width = vo_widths[c_idx]
            set_cell_margins(row.cells[c_idx], top=50, bottom=50, left=100, right=100)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # =========================================================================
    # 6. EXTRACTED DECLARATION EVIDENCE
    # =========================================================================
    h5 = doc.add_heading("5. Extracted Declaration Evidence", level=2)
    h5.paragraph_format.space_before = Pt(8)
    h5.paragraph_format.space_after = Pt(4)
    h5.paragraph_format.keep_with_next = True
    h5.runs[0].font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    if report.extracted_evidence:
        ev_table = doc.add_table(rows=len(report.extracted_evidence) + 1, cols=5)
        ev_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(ev_table)
        set_repeat_table_header(ev_table.rows[0])

        e_widths = [Inches(1.25), Inches(3.15), Inches(0.95), Inches(1.3), Inches(0.85)]
        e_headers = ["Declaration", "Final Detected Value", "Status", "Evidence Source", "Source View"]

        for c_idx, h_text in enumerate(e_headers):
            cell = ev_table.rows[0].cells[c_idx]
            cell.width = e_widths[c_idx]
            set_cell_background(cell, "F1F5F9")
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(h_text)
            run.font.size = Pt(8.5)
            run.bold = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

        for r_idx, ev in enumerate(report.extracted_evidence):
            row = ev_table.rows[r_idx + 1]
            view_str = ", ".join(ev.view_ids) if ev.view_ids else "—"

            row_vals = [
                _declaration_label(ev),
                _final_detected_value(ev),
                ev.status.replace("_", " ").title(),
                "Reconciled package evidence",
                view_str,
            ]
            for c_idx, text_val in enumerate(row_vals):
                cell = row.cells[c_idx]
                cell.width = e_widths[c_idx]
                set_cell_margins(cell, top=50, bottom=50, left=100, right=100)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(text_val)
                run.font.size = Pt(8.5)
                if c_idx == 0:
                    run.bold = True

        doc.add_paragraph().paragraph_format.space_after = Pt(4)
    else:
        p_no_ev = doc.add_paragraph()
        r = p_no_ev.add_run("No declaration candidates extracted from provided captures.")
        r.italic = True
        r.font.size = Pt(8.5)

    # =========================================================================
    # 6. TRACEABLE EVIDENCE IMAGERY & HIGHLIGHTED REGIONS (APPENDIX)
    # =========================================================================
    h6_ev = doc.add_heading("6. Traceable Evidence Imagery & Highlighted Regions", level=2)
    h6_ev.paragraph_format.space_before = Pt(8)
    h6_ev.paragraph_format.space_after = Pt(2)
    h6_ev.paragraph_format.keep_with_next = True
    h6_ev.runs[0].font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    ev_app_note = doc.add_paragraph()
    ev_app_note.paragraph_format.space_after = Pt(4)
    ev_app_note_run = ev_app_note.add_run(
        "Annotated optical crops illustrating detected declaration boundaries and statutory context. "
        "All annotations are derivative visual references; source image SHA-256 hashes are strictly preserved."
    )
    ev_app_note_run.italic = True
    ev_app_note_run.font.size = Pt(8.5)
    ev_app_note_run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    from app.services.image_store import get_capture_image, get_image_by_hash
    from app.services.evidence_annotation_service import annotate_evidence_crop

    evidence_cards_rendered = 0

    for ev in report.extracted_evidence:
        asset, exact_region_supported, use_polygons = resolve_evidence_asset(report, ev)
        img_bytes = None
        source_sha = asset.image_sha256 if asset else "—"
        source_cap_id = asset.capture_id if asset else "—"
        source_view = asset.view_id if asset else (", ".join(ev.view_ids) if ev.view_ids else "—")

        if asset:
            if getattr(asset, "image_b64", None):
                try:
                    import base64
                    img_bytes = base64.b64decode(asset.image_b64)
                except Exception:
                    img_bytes = None
            if not img_bytes:
                img_bytes = get_capture_image(asset.capture_id) or get_image_by_hash(asset.image_sha256)

        rendered_image = None
        evidence_label = "Contextual visual evidence — exact OCR region unavailable"
        if img_bytes and exact_region_supported:
            rendered_image = annotate_evidence_crop(
                image_bytes=img_bytes,
                field_label=ev.field,
                view_id=source_view,
                polygons=ev.polygons if use_polygons else [],
                pixel_box=ev.pixel_box,
                normalized_box=ev.normalized_box,
                is_clearance=(ev.field == "NET_QUANTITY"),
            )
            if rendered_image:
                evidence_label = "OCR-supported evidence region"
        if img_bytes and not rendered_image:
            rendered_image = img_bytes

        ev_id_str = ", ".join(ev.evidence_ids) if ev.evidence_ids else "—"
        sha_display = f"{source_sha[:16]}..." if source_sha != "—" else "—"
        confidence = f"{int(ev.confidence * 100)}%" if ev.confidence is not None else "—"

        if rendered_image:
            card_tbl = doc.add_table(rows=1, cols=2)
            card_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(card_tbl)

            c_img = card_tbl.rows[0].cells[0]
            c_img.width = Inches(3.2)
            set_cell_background(c_img, "F8FAFC")
            set_cell_margins(c_img, top=60, bottom=60, left=100, right=100)
            p_img = c_img.paragraphs[0]
            p_img.paragraph_format.space_after = Pt(0)
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_img = p_img.add_run()
            run_img.add_picture(io.BytesIO(rendered_image), width=Inches(3.0))

            c_txt = card_tbl.rows[0].cells[1]
            c_txt.width = Inches(4.3)
            set_cell_background(c_txt, "F8FAFC")
            set_cell_margins(c_txt, top=60, bottom=60, left=100, right=100)
            p_txt = c_txt.paragraphs[0]
            p_txt.paragraph_format.space_after = Pt(0)

            p_txt.add_run("Declaration: ").bold = True
            p_txt.add_run(f"{_declaration_label(ev)}\n")
            p_txt.add_run("Final Detected Value: ").bold = True
            p_txt.add_run(f"{_final_detected_value(ev)}\n")
            p_txt.add_run("Recorded Package Text: ").bold = True
            p_txt.add_run(f"{ev.raw_value or '—'}\n")
            p_txt.add_run("Machine Confidence: ").bold = True
            p_txt.add_run(f"{confidence}\n")
            p_txt.add_run("Evidence ID: ").bold = True
            r_eid = p_txt.add_run(f"{ev_id_str}\n")
            r_eid.font.name = "Courier New"
            r_eid.font.size = Pt(7.5)
            p_txt.add_run("Source Capture: ").bold = True
            p_txt.add_run(f"{source_view} ({source_cap_id[:10]}...)  |  ")
            p_txt.add_run("SHA: ").bold = True
            p_txt.add_run(f"{sha_display}\n")
            r_evidence_label = p_txt.add_run(evidence_label)
            r_evidence_label.italic = True
            r_evidence_label.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            evidence_cards_rendered += 1
        elif asset:
            fb_tbl = doc.add_table(rows=1, cols=1)
            fb_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(fb_tbl)
            c_fb = fb_tbl.rows[0].cells[0]
            c_fb.width = Inches(7.5)
            set_cell_background(c_fb, "F8FAFC")
            set_cell_margins(c_fb, top=60, bottom=60, left=100, right=100)
            p_fb = c_fb.paragraphs[0]
            p_fb.paragraph_format.space_after = Pt(0)
            p_fb.add_run(f"Declaration: {_declaration_label(ev)}  |  Final value: {_final_detected_value(ev)}  |  Evidence ID: {ev_id_str}\n").bold = True
            r_note = p_fb.add_run("Source image was unavailable for rendering.")
            r_note.italic = True
            r_note.font.color.rgb = RGBColor(0xD9, 0x77, 0x06)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            evidence_cards_rendered += 1

    if evidence_cards_rendered == 0:
        p_none = doc.add_paragraph()
        r_none = p_none.add_run("No visual evidence regions available for display.")
        r_none.italic = True
        r_none.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # =========================================================================
    # 7. INSPECTOR CONTEXT SNAPSHOT
    # =========================================================================
    h6 = doc.add_heading("7. Inspector-Provided Context Snapshot", level=2)
    h6.paragraph_format.space_before = Pt(8)
    h6.paragraph_format.space_after = Pt(2)
    h6.paragraph_format.keep_with_next = True
    h6.runs[0].font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    ctx_note = doc.add_paragraph()
    ctx_note.paragraph_format.space_after = Pt(4)
    ctx_note_run = ctx_note.add_run(
        "Explicit contextual parameters supplied by the inspector during progressive clarification. "
        "Unspecified parameters remain 'Unknown / Not established' without silent assumptions."
    )
    ctx_note_run.italic = True
    ctx_note_run.font.size = Pt(8.5)
    ctx_note_run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    ctx = report.inspector_context
    ctx_table = doc.add_table(rows=4, cols=4)
    ctx_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(ctx_table)

    ctx_col_widths = [Inches(1.4), Inches(2.35), Inches(1.4), Inches(2.35)]
    ctx_rows = [
        [("Product Origin Classification:", True), (_format_context_value(ctx.product_origin), False), ("Regulatory Class:", True), (_format_context_value(ctx.regulatory_product_class), False)],
        [("Date Regime:", True), (_format_context_value(ctx.date_regulatory_regime), False), ("Date Exemption:", True), (_format_context_value(ctx.date_package_exemption), False)],
        [("Electronic Commodity:", True), (_format_context_value(ctx.is_electronic), False), ("Package Structure:", True), (_format_context_value(ctx.package_structure), False)],
        [("Alcohol Context:", True), (_format_context_value(ctx.alcohol_context), False), ("", False), ("", False)]
    ]

    for r_idx, row_data in enumerate(ctx_rows):
        row = ctx_table.rows[r_idx]
        for c_idx, (text_val, is_bold) in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.width = ctx_col_widths[c_idx]
            set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text_val)
            run.font.size = Pt(8.5)
            run.bold = is_bold
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A) if is_bold else RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # =========================================================================
    # 8. STATUTORY LIMITATIONS & LEGAL DISCLAIMER
    # =========================================================================
    h7 = doc.add_heading("8. Statutory & System Limitations", level=2)
    h7.paragraph_format.space_before = Pt(8)
    h7.paragraph_format.space_after = Pt(4)
    h7.paragraph_format.keep_with_next = True
    h7.runs[0].font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    disclaimer_items = [
        ("1. Rule 7 Physical Numeral Height: ", "Uncalibrated 2D optical photographs cannot establish absolute millimeter physical dimensions without a certified millimeter calibration target and verified Principal Display Panel area measurement."),
        ("2. Rule 8 Principal Display Panel (PDP): ", "The image capture view role 'FRONT' does not automatically constitute the statutory Principal Display Panel as defined under Rule 2(h) of the Legal Metrology Rules, 2011."),
        ("3. Rule 9 Color Contrast: ", "Photographic ambient lighting and sensor white-balance vary; definitive background contrast evaluation requires standardized spectrophotometric measurement."),
        ("4. Decision Support Notice: ", "This inspection record is generated by DRISHTI for automated decision-support, statutory check verification, and evidence traceability. It does not constitute a judicial certificate or final adjudication order.")
    ]

    for lead, body in disclaimer_items:
        dp = doc.add_paragraph()
        dp.paragraph_format.space_after = Pt(2)
        r_lead = dp.add_run(lead)
        r_lead.bold = True
        r_lead.font.size = Pt(8.5)
        r_body = dp.add_run(body)
        r_body.font.size = Pt(8.5)
        r_body.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # Save to buffer
    buffer = io.BytesIO()
    doc.save(buffer)
    docx_bytes = buffer.getvalue()
    buffer.close()
    return docx_bytes
